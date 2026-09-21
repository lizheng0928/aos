"""Import Word content verbatim, retaining ordered source nodes and image bytes."""
import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run


def read_export(path, name):
    text = path.read_text()
    start = text.index("= [", text.index("export const " + name)) + 2
    return json.JSONDecoder().raw_decode(text[start:])[0]


def remove_card_label(block):
    if len(block["rows"]) != 1 or len(block["rows"][0]) != 1:
        return None
    cell = block["cells"][0][0]
    if not cell or cell[0]["type"] != "paragraph":
        return None
    paragraph = cell[0]
    original = "".join(r["text"] for r in paragraph["runs"])
    match = re.match(r"^[ \t]*(?:text|plain text)[ \t]*(?:\r?\n|$)", original, re.I)
    if not match:
        return None
    remaining = match.end()
    runs = []
    for run in paragraph["runs"]:
        skip = min(remaining, len(run["text"]))
        remaining -= skip
        if run["text"][skip:]:
            runs.append({**run, "text": run["text"][skip:]})
    paragraph["runs"] = runs
    if not runs:
        cell.pop(0)
    block["rows"][0][0] = re.sub(r"^[ \t]*(?:text|plain text)[ \t]*(?:\r?\n|$)", "", block["rows"][0][0], count=1, flags=re.I)
    return {"sourceId": paragraph["sourceId"], "originalText": original, "displayText": original[match.end():]}


def is_generated_footer(text):
    return text.strip().lstrip("|").strip() == "（注：部分内容可能由 AI 生成）"


def import_release(source, root):
    document = Document(source)
    data_path = root / "app/playbook-data.ts"
    old_chapters = read_export(data_path, "playbookChapters")
    old_ids = [s["id"] for c in old_chapters for s in c["sections"]]
    image_dir = root / "public/playbook-images"
    image_dir.mkdir(parents=True, exist_ok=True)
    numbering = document.part.numbering_part.element
    nums = {n.get(qn("w:numId")): n for n in numbering.findall(qn("w:num"))}
    abstracts = {n.get(qn("w:abstractNumId")): n for n in numbering.findall(qn("w:abstractNum"))}
    counters = Counter()
    chapters, sections, images, source_nodes = [], [], [], []
    removed_card_labels = []
    removed_source_notes = []
    image_names = {}
    current = None
    title = document.paragraphs[0].text

    def marker(paragraph):
        prop = paragraph._p.pPr
        if prop is None or prop.numPr is None or prop.numPr.numId is None:
            return {}
        num_id = str(prop.numPr.numId.val)
        level = prop.numPr.ilvl.val if prop.numPr.ilvl is not None else 0
        num = nums[num_id]
        abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        levels = abstracts[abstract_id].findall(qn("w:lvl"))
        definition = next(n for n in levels if int(n.get(qn("w:ilvl"), "0")) == level)
        fmt = definition.find(qn("w:numFmt")).get(qn("w:val"))
        template = definition.find(qn("w:lvlText")).get(qn("w:val"))
        if fmt == "bullet":
            value = template
        else:
            assert fmt == "decimal", f"Unsupported numbering: {fmt}"
            start_node = definition.find(qn("w:start"))
            start = int(start_node.get(qn("w:val"))) if start_node is not None else 1
            key = (num_id, level)
            value = template.replace(f"%{level + 1}", str(start + counters[key]))
            counters[key] += 1
        return {"marker": value, "level": level}

    def paragraph_blocks(paragraph, source_id, in_cell=False):
        result, runs = [], []
        list_meta = marker(paragraph)
        outline = paragraph._p.find("./" + qn("w:pPr") + "/" + qn("w:outlineLvl"))

        def flush():
            if not runs:
                return
            heading = outline is not None and not in_cell
            block = {"type": "heading" if heading else "paragraph", "runs": runs[:], "sourceId": source_id, **list_meta}
            if heading:
                block["headingLevel"] = min(6, int(outline.get(qn("w:val"))) + 1)
            result.append(block)
            runs.clear()

        for element in paragraph._p.iter():
            if element.tag != qn("w:r"):
                continue
            run = Run(element, paragraph)
            style = {}
            for key in ("bold", "italic", "underline"):
                if getattr(run, key):
                    style[key] = True
            if run.font.strike:
                style["strike"] = True
            parent = element.getparent()
            if parent.tag == qn("w:hyperlink"):
                rel_id = parent.get(qn("r:id"))
                if rel_id:
                    style["href"] = document.part.rels[rel_id].target_ref
            for child in element:
                if child.tag == qn("w:drawing"):
                    flush()
                    for blip in child.xpath(".//a:blip"):
                        rel_id = blip.get(qn("r:embed"))
                        part = document.part.related_parts[rel_id]
                        if rel_id not in image_names:
                            suffix = Path(str(part.partname)).suffix
                            image_names[rel_id] = f"release-image-{len(image_names)+1:03d}{suffix}"
                            (image_dir / image_names[rel_id]).write_bytes(part.blob)
                        extent = child.xpath(".//wp:extent")
                        props = child.xpath(".//wp:docPr")
                        image = {"type": "image", "src": "/playbook-images/" + image_names[rel_id], "alt": props[0].get("descr", "") if props else "", "sourceId": source_id}
                        if extent:
                            image.update(width=round(int(extent[0].get("cx")) / 9525), height=round(int(extent[0].get("cy")) / 9525))
                        result.append(image)
                        images.append({"sourceId": source_id, "src": image["src"], "sha256": hashlib.sha256(part.blob).hexdigest(), "relationship": rel_id})
                else:
                    value = None
                    if child.tag == qn("w:t"):
                        value = child.text or ""
                    elif child.tag in (qn("w:br"), qn("w:cr")):
                        value = "\n"
                    elif child.tag == qn("w:tab"):
                        value = "\t"
                    elif child.tag == qn("w:noBreakHyphen"):
                        value = "\u2011"
                    if value is not None:
                        runs.append({"text": value, **style})
        flush()
        actual = "".join(r["text"] for b in result if "runs" in b for r in b["runs"])
        assert actual == paragraph.text, (source_id, actual, paragraph.text)
        return result

    for index, child in enumerate(document.element.body):
        source_id = f"body-{index}"
        if child.tag == qn("w:p"):
            p = Paragraph(child, document)
            text = p.text
            if re.match(r"^Chapter [1-4]", text):
                chapters.append({"id": old_chapters[len(chapters)]["id"], "title": text, "sections": []})
            elif re.match(r"^[1-4]\.\d+\s", text):
                current = {"id": old_ids[len(sections)], "title": text, "chapterId": chapters[-1]["id"], "chapterTitle": chapters[-1]["title"], "index": len(sections)+1, "prev": None, "next": None, "summary": "", "blocks": []}
                sections.append(current)
                chapters[-1]["sections"].append({"id": current["id"], "title": text})
            elif current is not None:
                if is_generated_footer(text):
                    removed_source_notes.append({"sourceId": source_id, "originalText": text})
                    continue
                blocks = paragraph_blocks(p, source_id)
                current["blocks"].extend(blocks)
                source_nodes.extend({"section": current["id"], "block": b} for b in blocks)
        elif child.tag == qn("w:tbl"):
            assert current is not None
            table = Table(child, document)
            cells = [[[b for pi,p in enumerate(c.paragraphs) for b in paragraph_blocks(p, f"{source_id}/r{ri}/c{ci}/p{pi}", True)] for ci,c in enumerate(row.cells)] for ri,row in enumerate(table.rows)]
            rows = [[c.text for c in row.cells] for row in table.rows]
            block = {"type": "table", "rows": rows, "cells": cells, "sourceId": source_id}
            removed = remove_card_label(block)
            if removed:
                removed_card_labels.append(removed)
            # Only existing, explicitly identified diagrams replace source text.
            if len(rows) == 1 and len(rows[0]) == 1:
                text = rows[0][0]
                if text.startswith("State the goal →"):
                    block["diagram"] = "process"
                    block["stages"] = [part.strip() for part in text.split("→")]
                elif "Scenario A (Generate JD)" in text and "Candidate status updated" in text:
                    block["diagram"] = "hr"
                elif text.startswith("Receive goal → Observe materials & state"):
                    block["diagram"] = "process"
                    block["stages"] = [part.strip() for part in text.split("→")]
                elif "my-skill/" in text and "SKILL.md" in text:
                    block["diagram"] = "skill"
                elif "invoice-skill/" in text and "SKILL.md" in text:
                    block["diagram"] = "invoice"
            current["blocks"].append(block)
            source_nodes.append({"section": current["id"], "block": block})

    assert len(chapters) == 4 and len(sections) == 29
    assert len(images) == len(document.inline_shapes) == 119
    # The Word revision omits the diagram; the user explicitly retains this card.
    appendix = sections[-1]
    anchor = next(i for i,b in enumerate(appendix["blocks"]) if b.get("runs") and "".join(r["text"] for r in b["runs"]) == "The Big Picture: How the Agent Works")
    note = appendix["blocks"][anchor + 1]
    assert "".join(r["text"] for r in note["runs"]) == "The Agent loads Skills, calls tools, and sends results back to the checkpoint before final output."
    # This source paragraph is already the retained card's caption.
    appendix["blocks"][anchor + 1] = {**note, "type": "diagram", "diagram": "ai"}
    for i,section in enumerate(sections):
        section["prev"] = sections[i-1]["id"] if i else None
        section["next"] = sections[i+1]["id"] if i+1<len(sections) else None

    header = '''export type RichRun = { text: string; bold?: boolean; italic?: boolean; underline?: boolean; strike?: boolean; href?: string };
export type TextBlock = { type: 'paragraph' | 'heading'; runs: RichRun[]; marker?: string; level?: number; headingLevel?: number; sourceId: string };
export type ImageBlock = { type: 'image'; src: string; alt: string; width?: number; height?: number; sourceId: string };
export type PlaybookBlock = TextBlock | ImageBlock
  | { type: 'table'; rows: string[][]; cells: (TextBlock | ImageBlock)[][][]; diagram?: 'process' | 'hr' | 'skill' | 'invoice'; stages?: string[]; sourceId: string }
  | { type: 'diagram'; diagram: 'ai'; sourceId: string; runs: RichRun[] };
export type PlaybookSection = { id: string; title: string; chapterId: string; chapterTitle: string; index: number; prev: string | null; next: string | null; summary: string; blocks: PlaybookBlock[] };
export type PlaybookChapter = { id: string; title: string; sections: Array<{ id: string; title: string }> };
'''
    data_path.write_text(header + "\nexport const playbookTitle = " + json.dumps(title, ensure_ascii=False) + ";\nexport const playbookChapters: PlaybookChapter[] = " + json.dumps(chapters, ensure_ascii=False, indent=2) + ";\n\nexport const playbookSections: PlaybookSection[] = " + json.dumps(sections, ensure_ascii=False, indent=2) + ';\n\nexport function getSectionById(id: string) {\n  return playbookSections.find((section) => section.id === id);\n}\n')
    audit_dir = root / "qa"
    audit_dir.mkdir(exist_ok=True)
    (audit_dir / "source-manifest.json").write_text(json.dumps({"sourceFile": source.name, "sourceSha256": hashlib.sha256(source.read_bytes()).hexdigest(), "chapters": chapters, "sections": sections, "images": images, "sourceNodes": source_nodes, "removedCardLabels": removed_card_labels, "removedSourceNotes": removed_source_notes}, ensure_ascii=False, indent=2))
    print(json.dumps({"chapters": len(chapters), "sections": len(sections), "tables": len(document.tables), "imageOccurrences": len(images), "uniqueImages": len(image_names), "sourceNodes": len(source_nodes)}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    import_release(args.source, args.root)
